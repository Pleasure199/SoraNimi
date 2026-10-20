"""Bounded text extraction; binary document parsers run in a separate process."""

import io
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 500_000
MAX_PAGES = 500
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".log", ".tsv"}
DOCUMENT_SUFFIXES = {".pdf", ".docx", ".pptx"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | DOCUMENT_SUFFIXES


class DocumentError(ValueError):
    pass


def validate_filename(filename):
    if (
        not filename
        or len(filename) > 200
        or any(c in filename for c in "/\\")
        or any(ord(c) < 32 for c in filename)
    ):
        raise DocumentError("文件名无效（不能包含路径或控制字符，最长 200 字符）")
    suffix = Path(filename).suffix.lower()
    if suffix in {".doc", ".ppt"}:
        raise DocumentError("旧版 .doc/.ppt 请先另存为 .docx/.pptx 后上传")
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentError("暂不支持此格式；支持 TXT/MD/CSV/JSON/LOG/TSV、PDF、DOCX、PPTX")
    return suffix


def parse_document(data, filename):
    suffix = validate_filename(filename)
    if not data:
        raise DocumentError("文件为空")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError("文件超过 20 MiB 限制")
    if suffix in TEXT_SUFFIXES:
        from soranimi.tools.read_file import decode_text

        try:
            text, encoding = decode_text(data)
        except UnicodeError:
            try:
                text, encoding = data.decode("gb18030"), "gb18030"
            except UnicodeError as exc:
                raise DocumentError("无法识别文本编码，请转为 UTF-8 后上传") from exc
        if "\x00" in text:
            raise DocumentError("文件包含二进制内容，不能作为文本解析")
        return result(text, encoding=encoding)
    # No provider credentials, external commands, macros or network fetching.
    env = {
        k: v
        for k, v in os.environ.items()
        if k.upper() in {"PATH", "SYSTEMROOT", "WINDIR", "TEMP", "TMP"}
    }
    env["PYTHONIOENCODING"] = "utf-8"
    try:
        process = subprocess.run(
            [sys.executable, "-m", "soranimi.documents", suffix],
            input=data,
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
            timeout=30,
            env=env,
        )
    except subprocess.TimeoutExpired as exc:
        raise DocumentError("文档解析超时，请拆分文档后重试") from exc
    if process.returncode or len(process.stdout) > 8 * MAX_TEXT_CHARS:
        raise DocumentError("文档解析失败或超出资源限制")
    try:
        output = json.loads(process.stdout)
    except (UnicodeError, ValueError) as exc:
        raise DocumentError("文档解析失败") from exc
    if "error" in output:
        raise DocumentError(output["error"])
    return output


def result(text, **metadata):
    if not text.strip():
        raise DocumentError("未提取到文本；扫描件或纯图片文档需要先进行 OCR")
    # Wrap long paragraphs so every retained character is reachable by line paging.
    clipped = text[:MAX_TEXT_CHARS]
    wrapped = "\n".join(
        chunk
        for line in clipped.splitlines()
        for chunk in ([line[i : i + 2000] for i in range(0, len(line), 2000)] or [""])
    )
    return {
        "text": wrapped[:MAX_TEXT_CHARS],
        "truncated": len(text) > MAX_TEXT_CHARS or len(wrapped) > MAX_TEXT_CHARS,
        **metadata,
    }


def check_office_archive(data):
    from defusedxml.ElementTree import fromstring

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        items = archive.infolist()
        if len(items) > 3000 or sum(item.file_size for item in items) > 80 * 1024 * 1024:
            raise DocumentError("Office 文档解压后过大或包含过多条目")
        for item in items:
            if item.flag_bits & 1:
                raise DocumentError("不支持加密 Office 文档，请解密后上传")
            if item.filename.lower().endswith((".xml", ".rels")):
                # Reject entities/DTDs, including UTF-16 encoded XML, before lxml loads it.
                fromstring(archive.read(item), forbid_dtd=True)


def extract_binary(data, suffix):
    parts = []
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError("不支持加密 PDF，请解密后上传")
        if len(reader.pages) > MAX_PAGES:
            raise DocumentError("PDF 超过 500 页限制，请拆分后上传")
        for index, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                parts.append(f"[Page {index}]\n{text}")
            if sum(map(len, parts)) > MAX_TEXT_CHARS:
                break
        return result("\n\n".join(parts), pages=len(reader.pages))
    check_office_archive(data)
    if suffix == ".docx":
        from docx import Document
        from docx.table import Table

        document = Document(io.BytesIO(data))
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                parts.append(
                    "\n".join("\t".join(cell.text for cell in row.cells) for row in block.rows)
                )
            else:
                parts.append(block.text)
            if sum(map(len, parts)) > MAX_TEXT_CHARS:
                break
        return result("\n".join(parts))
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(io.BytesIO(data))
    if len(presentation.slides) > MAX_PAGES:
        raise DocumentError("PPT 超过 500 张幻灯片限制，请拆分后上传")

    def shape_text(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shape_text(shape.shapes)
            elif shape.has_text_frame:
                yield shape.text
            elif shape.has_table:
                for row in shape.table.rows:
                    yield "\t".join(cell.text for cell in row.cells)

    for index, slide in enumerate(presentation.slides, 1):
        texts = list(shape_text(slide.shapes))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            texts.append("[Notes]\n" + slide.notes_slide.notes_text_frame.text)
        text = "\n".join(texts)
        if text.strip():
            parts.append(f"[Slide {index}]\n{text}")
        if sum(map(len, parts)) > MAX_TEXT_CHARS:
            break
    return result("\n\n".join(parts), slides=len(presentation.slides))


def main():
    # Process/resource limits are not an OS security sandbox for arbitrary code.
    if os.name != "nt":
        import resource

        resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
        try:
            resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
        except (ValueError, OSError):
            pass
    try:
        data = sys.stdin.buffer.read(MAX_FILE_BYTES + 1)
        if len(data) > MAX_FILE_BYTES:
            raise DocumentError("文件超过 20 MiB 限制")
        output = extract_binary(data, sys.argv[1])
    except ImportError:
        output = {"error": "缺少文档解析依赖，请重新安装项目：pip install -e ."}
    except DocumentError as exc:
        output = {"error": str(exc)}
    except Exception:
        output = {"error": "文档损坏、格式不匹配或无法安全解析"}
    sys.stdout.write(json.dumps(output, ensure_ascii=False))


if __name__ == "__main__":
    main()
