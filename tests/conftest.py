from collections.abc import Sequence

import pytest
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.messages import AIMessage, BaseMessage
from langchain_core.outputs import ChatGeneration, ChatResult
from pydantic import Field

from soranimi.agent.demo import DemoModel
from soranimi.config import Settings
from soranimi.runtime import AgentSession


class ScriptedModel(BaseChatModel):
    responses: Sequence[AIMessage]
    cursor: int = 0
    observed: list[list[BaseMessage]] = Field(default_factory=list)

    @property
    def _llm_type(self):
        return "test-scripted"

    def bind_tools(self, tools, **kwargs):
        return self

    def _generate(self, messages, stop=None, run_manager=None, **kwargs):
        self.observed.append(list(messages))
        message = self.responses[self.cursor]
        self.cursor += 1
        return ChatResult(generations=[ChatGeneration(message=message)])


def calls(*items):
    return AIMessage(
        content="",
        tool_calls=[
            {"name": name, "args": args, "id": f"call-{i}"} for i, (name, args) in enumerate(items)
        ],
    )


@pytest.fixture
def session(tmp_path):
    with AgentSession(Settings(workspace=tmp_path), DemoModel()) as value:
        yield value


@pytest.fixture
def tools(session):
    return {tool.name: tool for tool in session.tools}


def document_bytes(suffix):
    """Small real documents; no network, Office installation or paid models."""
    from io import BytesIO

    stream = BytesIO()
    if suffix == ".txt":
        return "测试资料\nBudget: 42\n".encode()
    if suffix == ".docx":
        from docx import Document

        doc = Document()
        doc.add_paragraph("测试资料")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Budget"
        table.cell(0, 1).text = "42"
        doc.add_paragraph("After table")
        doc.save(stream)
    elif suffix == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches

        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "测试资料"
        table = slide.shapes.add_table(1, 2, 0, Inches(1), Inches(4), Inches(1)).table
        table.cell(0, 0).text = "Budget"
        table.cell(0, 1).text = "42"
        slide.notes_slide.notes_text_frame.text = "Speaker note"
        doc.save(stream)
    elif suffix == ".pdf":
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        doc = PdfWriter()
        page = doc.add_blank_page(width=600, height=800)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): doc._add_object(font)})}
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 10 700 Td (Budget: 42) Tj ET")
        page[NameObject("/Contents")] = doc._add_object(content)
        doc.write(stream)
    return stream.getvalue()
